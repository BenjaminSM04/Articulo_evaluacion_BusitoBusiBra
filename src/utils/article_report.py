"""Spanish article/report generation helpers for BUSI -> BUS-BRA experiments."""
# ruff: noqa: E501
from __future__ import annotations

import platform
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from ..evaluation.artifacts import METRIC_COLUMNS


def _fmt(value: Any) -> str:
    if value is None or value == "":
        return "NA"
    try:
        if pd.isna(value):
            return "NA"
    except TypeError:
        pass
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)


def _table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if df is None or df.empty:
        return "_NA: no hay resultados disponibles para esta sección._"
    show = df.copy()
    if max_rows is not None:
        show = show.head(max_rows)
    return show.to_markdown(index=False)


def write_environment_report(path: Path, extra: dict[str, Any] | None = None) -> Path:
    """Write environment provenance used by the article run."""
    extra = extra or {}
    lines = [
        "# Entorno de ejecución",
        "",
        f"- Fecha y hora: {datetime.now().isoformat(timespec='seconds')}",
        f"- Sistema operativo: {platform.platform()}",
        f"- Python: {platform.python_version()}",
    ]
    try:
        import torch
        import torchvision

        lines.extend([
            f"- PyTorch: {torch.__version__}",
            f"- torchvision: {torchvision.__version__}",
            f"- CUDA disponible: {torch.cuda.is_available()}",
            f"- CUDA PyTorch: {torch.version.cuda}",
            "- GPU: "
            + (torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU"),
        ])
    except Exception as exc:  # pragma: no cover - defensive provenance only
        lines.append(f"- PyTorch/CUDA: no disponible ({exc})")

    for module_name in ("sklearn", "pandas", "numpy", "matplotlib"):
        try:
            mod = __import__(module_name)
            lines.append(f"- {module_name}: {getattr(mod, '__version__', 'NA')}")
        except Exception:
            lines.append(f"- {module_name}: NA")

    for key, value in extra.items():
        lines.append(f"- {key}: {value}")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def summarize_metrics(metrics_df: pd.DataFrame, group_cols: list[str]) -> pd.DataFrame:
    if metrics_df.empty:
        return pd.DataFrame()
    if any(col not in metrics_df.columns for col in group_cols):
        return pd.DataFrame()
    cols = [c for c in METRIC_COLUMNS if c in metrics_df.columns]
    rows = []
    for keys, group in metrics_df.groupby(group_cols, dropna=False):
        keys = keys if isinstance(keys, tuple) else (keys,)
        row = dict(zip(group_cols, keys, strict=False))
        row["n_runs"] = int(len(group))
        for col in cols:
            vals = pd.to_numeric(group[col], errors="coerce")
            row[f"{col}_mean"] = float(vals.mean()) if vals.notna().any() else float("nan")
            row[f"{col}_std"] = float(vals.std(ddof=0)) if vals.notna().any() else float("nan")
        rows.append(row)
    return pd.DataFrame(rows)


def compute_drop_table(metrics_df: pd.DataFrame, source: str, target: str) -> pd.DataFrame:
    """Compute internal BUSI vs external BUS-BRA absolute and percentage drops."""
    if metrics_df.empty:
        return pd.DataFrame()
    rows = []
    for arch in sorted(metrics_df.get("arch", pd.Series(dtype=str)).dropna().unique()):
        internal = metrics_df[
            (metrics_df.get("phase") == "baseline_internal")
            & (metrics_df.get("dataset") == source)
            & (metrics_df.get("arch") == arch)
        ]
        external = metrics_df[
            (metrics_df.get("phase") == "external_all")
            & (metrics_df.get("source") == source)
            & (metrics_df.get("target") == target)
            & (metrics_df.get("arch") == arch)
        ]
        if internal.empty or external.empty:
            continue
        for metric in METRIC_COLUMNS:
            if metric not in internal.columns or metric not in external.columns:
                continue
            i_val = pd.to_numeric(internal[metric], errors="coerce").mean()
            e_val = pd.to_numeric(external[metric], errors="coerce").mean()
            if pd.isna(i_val) or pd.isna(e_val):
                continue
            drop = i_val - e_val
            pct = (drop / i_val * 100.0) if i_val != 0 else float("nan")
            rows.append({
                "arch": arch,
                "metric": metric,
                "internal_busi": i_val,
                "external_busbra": e_val,
                "absolute_drop": drop,
                "percent_drop": pct,
            })
    return pd.DataFrame(rows)


def compute_recovery_table(metrics_df: pd.DataFrame, source: str, target: str) -> pd.DataFrame:
    """Compare adaptation/fine-tuning methods against the direct external baseline."""
    if metrics_df.empty:
        return pd.DataFrame()
    rows = []
    for arch in sorted(metrics_df.get("arch", pd.Series(dtype=str)).dropna().unique()):
        base = metrics_df[
            (metrics_df.get("phase") == "adaptation")
            & (metrics_df.get("method") == "none")
            & (metrics_df.get("source") == source)
            & (metrics_df.get("target") == target)
            & (metrics_df.get("arch") == arch)
        ]
        if base.empty:
            base = metrics_df[
                (metrics_df.get("phase") == "external_test")
                & (metrics_df.get("source") == source)
                & (metrics_df.get("target") == target)
                & (metrics_df.get("arch") == arch)
            ]
        if base.empty:
            continue
        adapted = metrics_df[
            (metrics_df.get("phase").isin(["adaptation", "finetuning"]))
            & (metrics_df.get("source") == source)
            & (metrics_df.get("target") == target)
            & (metrics_df.get("arch") == arch)
        ]
        for _, row in adapted.iterrows():
            method = row.get("method", row.get("phase", "NA"))
            for metric in METRIC_COLUMNS:
                if metric not in row or metric not in base.columns:
                    continue
                base_val = pd.to_numeric(base[metric], errors="coerce").mean()
                cur_val = pd.to_numeric(pd.Series([row[metric]]), errors="coerce").iloc[0]
                if pd.isna(base_val) or pd.isna(cur_val):
                    continue
                improvement = (
                    base_val - cur_val if metric in {"brier", "ece"} else cur_val - base_val
                )
                rows.append({
                    "arch": arch,
                    "method": method,
                    "metric": metric,
                    "external_none": base_val,
                    "method_value": cur_val,
                    "recovery_or_improvement": improvement,
                })
    return pd.DataFrame(rows)


def _best_method_text(recovery_df: pd.DataFrame) -> str:
    if recovery_df.empty or "metric" not in recovery_df.columns:
        return "NA: no se completaron comparaciones de recuperación."
    auc_rows = recovery_df[recovery_df["metric"] == "auc"].copy()
    if auc_rows.empty:
        return "NA: no hay AUC disponible para comparar recuperación."
    best = auc_rows.sort_values("recovery_or_improvement", ascending=False).iloc[0]
    return (
        f"{best['method']} en {best['arch']} "
        f"(mejora de AUC={_fmt(best['recovery_or_improvement'])})."
    )


def _write_docx(markdown_text: str, path: Path) -> bool:
    try:
        from docx import Document
    except Exception:
        return False
    doc = Document()
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            doc.add_paragraph(stripped)
    doc.save(path)
    return True


def _write_pdf(markdown_text: str, path: Path) -> bool:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except Exception:
        return False
    styles = getSampleStyleSheet()
    story = []
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue
        style = styles["Heading1"] if stripped.startswith("# ") else styles["BodyText"]
        text = stripped.lstrip("#").strip()
        story.append(Paragraph(text.replace("|", " | "), style))
        story.append(Spacer(1, 4))
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    doc.build(story)
    return True


def write_article_reports(
    report_dir: Path,
    dataset_summary: pd.DataFrame,
    metrics_df: pd.DataFrame,
    drop_df: pd.DataFrame,
    recovery_df: pd.DataFrame,
    status_df: pd.DataFrame,
    gradcam_summary: pd.DataFrame | None = None,
) -> dict[str, Path | None]:
    """Generate comparative_report.{md,docx,pdf} and article_materials.md."""
    report_dir.mkdir(parents=True, exist_ok=True)
    phase = metrics_df["phase"] if "phase" in metrics_df.columns else pd.Series([], dtype=str)
    baseline_table = summarize_metrics(
        metrics_df[phase == "baseline_internal"]
        if not metrics_df.empty and len(phase)
        else metrics_df,
        ["dataset", "arch"],
    )
    external_table = summarize_metrics(
        metrics_df[phase.isin(["external_all", "external_test"])]
        if not metrics_df.empty and len(phase) else metrics_df,
        ["phase", "arch"],
    )
    adaptation_table = summarize_metrics(
        metrics_df[phase == "adaptation"] if not metrics_df.empty and len(phase) else metrics_df,
        ["arch", "method"],
    )
    finetune_table = summarize_metrics(
        metrics_df[phase == "finetuning"] if not metrics_df.empty and len(phase) else metrics_df,
        ["arch", "method"],
    )

    best_method = _best_method_text(recovery_df)
    gradcam_text = (
        _table(gradcam_summary) if gradcam_summary is not None and not gradcam_summary.empty
        else "NA: Grad-CAM no fue ejecutado o no produjo métricas cuantitativas."
    )

    md = f"""# Generalización externa y adaptación de dominio en la clasificación de lesiones mamarias por ecografía mediante deep learning: evaluación BUSI -> BUS-BRA

## 1. Resumen ejecutivo
Se ejecutó o preparó una evaluación reproducible de clasificación binaria benigno/maligno usando BUSI como fuente y BUS-BRA como dominio externo. No se incluye la clase `normal` de BUSI en la comparación principal.

Mejor recuperación observada: {best_method}

## 2. Objetivo del experimento
Medir la generalización externa de modelos entrenados en BUSI y evaluados en BUS-BRA, cuantificar la caída de rendimiento y comparar adaptación de dominio, fine-tuning y explicabilidad Grad-CAM.

## 3. Datasets utilizados
{_table(dataset_summary)}

## 4. Metodología
Arquitecturas: ResNet-50, ResNet-18, EfficientNet-B0 y DenseNet-121 cuando sus experimentos están disponibles. El preprocesamiento redimensiona imágenes, armoniza etiquetas benigno/maligno, excluye `normal` y conserva máscaras solo para explicabilidad. Las particiones BUS-BRA de adaptación/test son agrupadas por paciente cuando hay `patient_id`.

## 5. Resultados internos en BUSI
{_table(baseline_table)}

## 6. Resultados externos en BUS-BRA
{_table(external_table)}

## 7. Caída de rendimiento BUSI -> BUS-BRA
{_table(drop_df)}

## 8. Resultados con adaptación de dominio
{_table(adaptation_table)}

## 9. Resultados con fine-tuning
{_table(finetune_table)}

## 10. Explicabilidad con Grad-CAM
{gradcam_text}

## 11. Calibración
Brier Score y ECE se reportan junto con las curvas de calibración cuando hay predicciones disponibles. En Brier/ECE, valores menores indican mejor calibración.

## 12. Discusión preliminar
La discusión debe apoyarse solo en las tablas generadas. Si una celda aparece como NA, el experimento no produjo evidencia suficiente y no debe interpretarse como resultado.

## 13. Conclusiones preliminares
Las conclusiones finales deben redactarse después de revisar la tabla de caída, la recuperación por método y Grad-CAM.

## 14. Material útil para artículo científico
Pregunta de investigación: ¿cuánto se degrada la clasificación benigno/maligno al pasar de BUSI a BUS-BRA y qué estrategias recuperan rendimiento sin comprometer la atención a la lesión?

Hipótesis: el cambio de dominio reduce el rendimiento externo y una adaptación/fine-tuning controlado recupera parte de la caída.

Tablas recomendadas: descripción de datasets, baselines internos, validación externa, caída de rendimiento, adaptación/fine-tuning, calibración y Grad-CAM.

Figuras recomendadas: matrices de confusión, curvas ROC, curvas PR, curvas de calibración, heatmaps de Grad-CAM.

## Estado de ejecución
{_table(status_df)}
"""

    md_path = report_dir / "comparative_report.md"
    md_path.write_text(md, encoding="utf-8")

    article = """# Material para artículo científico

## A. Título tentativo
Generalización externa y adaptación de dominio en la clasificación de lesiones mamarias por ecografía mediante deep learning: evaluación BUSI -> BUS-BRA

## B. Planteamiento del problema
Los modelos entrenados en un dataset pueden perder rendimiento al evaluarse en otra población o dispositivo.

## C. Justificación
Cuantificar el cambio de dominio es necesario antes de proponer uso clínico o investigación translacional.

## D. Pregunta de investigación
¿Qué tan bien generalizan modelos entrenados en BUSI cuando se evalúan en BUS-BRA y qué estrategias recuperan rendimiento?

## E. Objetivo general
Evaluar generalización externa, adaptación de dominio, fine-tuning y explicabilidad en clasificación benigno/maligno.

## F. Objetivos específicos
- Entrenar modelos internos en BUSI.
- Evaluar validación externa BUSI -> BUS-BRA sin reentrenar.
- Comparar adaptación de dominio y fine-tuning supervisado.
- Evaluar calibración y Grad-CAM.

## G. Hipótesis
El rendimiento cae bajo domain shift y la adaptación controlada recupera parte de esa caída.

## H. Metodología resumida
Clasificación binaria, exclusión de `normal`, splits reproducibles, métricas discriminativas, calibración y XAI con máscaras si existen.

## I. Resultados principales
Ver `comparative_report.md`; no completar manualmente con valores no generados.

## J. Discusión preliminar
Interpretar caída, recuperación, calibración y plausibilidad anatómica de Grad-CAM.

## K. Limitaciones
Tamaño muestral, sesgo de datasets públicos, posibles duplicados, diferencias de adquisición, y costo computacional.

## L. Conclusión preliminar
La utilidad científica depende de si la adaptación mejora el rendimiento externo y mantiene atención sobre la lesión.

## M. Tablas recomendadas
Datasets, métricas internas, métricas externas, caída, adaptación, fine-tuning, calibración, Grad-CAM.

## N. Figuras recomendadas
ROC, PR, calibración, matrices de confusión, heatmaps Grad-CAM y ejemplos por clase.
"""
    article_path = report_dir / "article_materials.md"
    article_path.write_text(article, encoding="utf-8")

    docx_path = report_dir / "comparative_report.docx"
    pdf_path = report_dir / "comparative_report.pdf"
    docx_ok = _write_docx(md, docx_path)
    pdf_ok = _write_pdf(md, pdf_path)

    return {
        "markdown": md_path,
        "article_materials": article_path,
        "docx": docx_path if docx_ok else None,
        "pdf": pdf_path if pdf_ok else None,
    }
